"""Integration coverage for LangGraph / API; requires full requirements installation."""
import io
import uuid
from docx import Document
from fastapi.testclient import TestClient
from app.main import app
from app.schemas import InterviewSummary
from app.services.documents import read_transcript


def send(client, text, app_id=None, file_id=None, msg_id=None):
    return client.post('/api/v1/messages',json={
        'source':'web','external_message_id':msg_id or uuid.uuid4().hex,
        'sender_id':'hr-demo','application_id':app_id,
        'file_id':file_id,'content':text
    })


def test_end_to_end_upload_rounds_and_review():
    with TestClient(app) as client:
        document=Document()
        for line in ['姓名：张晨','岗位：Python后端工程师','技能：FastAPI、SQL','学历：本科']:
            document.add_paragraph(line)
        buffer=io.BytesIO();document.save(buffer)
        up=client.post('/api/v1/files', files={'file':('zhangchen.docx',buffer.getvalue(),
                                                         'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        assert up.status_code==201,up.text
        assessed=client.post('/api/v1/resume-assessments',json={'file_id':up.json()['id'],'ai_config':{'mode':'demo'}})
        assert assessed.status_code==201,assessed.text
        created=client.post(f"/api/v1/resume-assessments/{assessed.json()['id']}/decision",json={
            'decision':'PENDING','evaluation':assessed.json()['evaluation']})
        assert created.status_code==200,created.text
        aid=created.json()['application_id']
        assert created.json()['new_stage']=='SCREENING'
        initial=send(client,'HR初筛评价：经验匹配，建议面试',aid)
        assert initial.json()['new_stage']=='ROUND1_PENDING'
        scheduled=send(client,'安排一面，2026-11-02 15:00，面试官李工',aid)
        assert scheduled.json()['new_stage']=='ROUND1_SCHEDULED'
        reviewed=send(client,'一面结束，沟通良好，一面通过，安排二面',aid)
        assert reviewed.json()['new_stage']=='ROUND2_PENDING'
        invalid=send(client,'安排二面，面试官王工',aid)
        assert invalid.json()['status']=='pending_review'
        rid=client.get('/api/v1/reviews').json()[0]['id']
        approval=client.post(f'/api/v1/reviews/{rid}/resolve',json={'action':'approve','scheduled_at':'2026-11-04 14:00'})
        assert approval.status_code==200,approval.text
        assert approval.json()['new_stage']=='ROUND2_SCHEDULED'
        detail=client.get('/api/v1/applications/'+aid).json()
        assert len(detail['interviews'])==2
        assert len(detail['events'])==5
        dup_id='duplicate-key'
        once=send(client,'二面结束，二面通过，安排三面',aid,msg_id=dup_id)
        twice=send(client,'二面结束，二面通过，安排三面',aid,msg_id=dup_id)
        assert once.json()['new_stage']=='ROUND3_PENDING'
        assert twice.json()['status']=='duplicate'


def test_unknown_review_allows_human_to_choose_and_complete_intent():
    from app.services.demo_seed import demo_id

    with TestClient(app) as client:
        application_id=demo_id('application:gu')
        pending=send(client,'这是一条需要人工判断的消息',application_id)
        assert pending.json()['status']=='pending_review'
        assert pending.json()['proposal']['intent']=='UNKNOWN'
        review_id=client.get('/api/v1/reviews').json()[0]['id']
        url=f'/api/v1/reviews/{review_id}/resolve'
        missing=client.post(url,json={'action':'approve','intent':'INTERVIEW_FEEDBACK',
                                      'application_id':application_id,'round_no':1})
        assert missing.status_code==422
        assert '面试评价不能为空' in missing.json()['detail']
        corrected=client.post(url,json={'action':'approve','intent':'INTERVIEW_FEEDBACK',
                                        'application_id':application_id,'round_no':1,
                                        'evaluation':'项目经验清晰，待进一步确认',
                                        'decision':'PENDING'})
        assert corrected.status_code==200,corrected.text
        assert corrected.json()['new_stage']=='ROUND1_FEEDBACK'


def test_model_connection_endpoint_reports_success_and_redacts_errors(monkeypatch, caplog):
    from app.api import model as model_api
    seen=[]
    monkeypatch.setattr(model_api.model,'check_connection',lambda config: seen.append(config))
    config={'mode':'custom','model':'test-model','base_url':'https://example.com/v1','api_key':'secret-test-key'}
    with TestClient(app) as client:
        response=client.post('/api/v1/model/test',json=config)
        assert response.status_code==200 and response.json()['model']=='test-model'
        assert seen[0].api_key=='secret-test-key'
        assert 'secret-test-key' not in response.text
        assert client.post('/api/v1/model/test',json={'mode':'demo'}).status_code==422
        def fail(_): raise RuntimeError('secret-test-key')
        monkeypatch.setattr(model_api.model,'check_connection',fail)
        failed=client.post('/api/v1/model/test',json=config)
        assert failed.status_code==502
        assert 'secret-test-key' not in failed.text and 'secret-test-key' not in caplog.text


def test_review_transaction_rolls_back_and_retries_after_failure(monkeypatch):
    from app.api import reviews as api
    from app.services.demo_seed import demo_id
    from app.db import SessionLocal
    from app.entities import Event
    from sqlalchemy import select

    with TestClient(app) as client:
        aid=demo_id('application:lin')
        pending=send(client,'暂时无法识别',aid).json()
        review_id=client.get('/api/v1/reviews').json()[0]['id']
        path=f'/api/v1/reviews/{review_id}/resolve'
        payload={'action':'approve','intent':'HR_EVALUATION','application_id':aid,'evaluation':'建议面试'}
        original=api.apply_proposal
        def fail_after_write(*args,**kwargs):
            original(*args,**kwargs)
            raise RuntimeError('模拟事务失败')
        monkeypatch.setattr(api,'apply_proposal',fail_after_write)
        assert client.post(path,json=payload).status_code==409
        assert client.get(f'/api/v1/applications/{aid}').json()['application']['stage']=='SCREENING'
        assert client.get('/api/v1/reviews').json()[0]['status']=='pending'
        with SessionLocal() as db:
            assert not db.scalars(select(Event).where(Event.message_id==pending['message_id'])).all()
        monkeypatch.setattr(api,'apply_proposal',original)
        retried=client.post(path,json=payload)
        assert retried.status_code==200,retried.text
        assert retried.json()['new_stage']=='ROUND1_PENDING'
        assert client.get('/api/v1/reviews').json()[0]['status']=='approved'


def test_review_checkpoint_failure_is_recoverable(monkeypatch):
    from app.api import reviews as api
    from app.services.demo_seed import demo_id
    from app.services.review_checkpoints import recover_review_checkpoints

    with TestClient(app) as client:
        aid=demo_id('application:lin')
        pending=send(client,'无法识别的消息',aid).json()
        review_id=client.get('/api/v1/reviews').json()[0]['id']
        original=api.resume_review_checkpoint
        def fail_checkpoint(*args): raise RuntimeError('模拟检查点失败')
        monkeypatch.setattr(api,'resume_review_checkpoint',fail_checkpoint)
        response=client.post(f'/api/v1/reviews/{review_id}/resolve',json={
            'action':'approve','intent':'HR_EVALUATION','application_id':aid,'evaluation':'建议面试'})
        assert response.status_code==200
        assert client.get('/api/v1/reviews').json()[0]['status']=='approved'
        thread={'configurable':{'thread_id':pending['message_id']}}
        assert app.state.graph.get_state(thread).next
        monkeypatch.setattr(api,'resume_review_checkpoint',original)
        recover_review_checkpoints(app.state.graph)
        assert not app.state.graph.get_state(thread).next


def test_interview_transcript_upload_and_summary(monkeypatch):
    from app.api import interviews as api
    captured = {}

    def fake_summary(text, candidate, position, round_no, config=None):
        captured.update(text=text, candidate=candidate, position=position, round_no=round_no,
                        mode=config.mode, model=config.model)
        return InterviewSummary(candidate_name='张晨', position=position, round_no=round_no,
                                interviewer='李工', interview_time='2026-11-02 15:00',
                                summary='讨论了 Python 基础与项目经验。', topics=['Python 基础'],
                                strengths=['解释了项目中的缓存策略'])

    monkeypatch.setattr(api, 'summarize_interview', fake_summary)
    with TestClient(app) as client:
        resume=send(client, '姓名：张晨，岗位：Python后端工程师，这是简历')
        aid=resume.json()['application_id']
        send(client, 'HR初筛评价：建议面试', aid)
        send(client, '安排一面，2026-11-02 15:00，面试官李工', aid)
        interview=client.get('/api/v1/applications/'+aid).json()['interviews'][0]
        transcript='面试官：请介绍项目。\n张晨：我实现了缓存策略。'
        response=client.post(f"/api/v1/interviews/{interview['id']}/transcripts",
                             data={'ai_mode':'custom','ai_model':'test-model',
                                   'ai_base_url':'https://example.com/v1','ai_api_key':'private-test-key'},
                             files={'file':('interview.txt', transcript.encode('utf-8'), 'text/plain')})
        assert response.status_code==201, response.text
        assert captured == {'text':transcript,'candidate':'张晨',
                            'position':'Python后端工程师','round_no':1,
                            'mode':'custom','model':'test-model'}
        detail=client.get('/api/v1/applications/'+aid).json()
        assert detail['interviews'][0]['transcript']['text']==transcript
        assert detail['interviews'][0]['transcript']['model']=='test-model'
        assert detail['interviews'][0]['transcript']['summary']['topics']==['Python 基础']
        assert 'private-test-key' not in str(detail)
        assert detail['application']['stage']=='ROUND1_SCHEDULED'
        assert client.post(f"/api/v1/interviews/{interview['id']}/transcripts",
                           files={'file':('interview.csv',b'a,b','text/csv')}).status_code==400
        monkeypatch.setattr(api, 'summarize_interview', lambda *args: InterviewSummary(
            candidate_name='其他人', summary='不应保存到张晨的记录'))
        mismatch=client.post(f"/api/v1/interviews/{interview['id']}/transcripts",
                             files={'file':('wrong.txt',b'wrong interview','text/plain')})
        assert mismatch.status_code==201
        decision=client.post(f"/api/v1/interviews/{interview['id']}/transcripts/{mismatch.json()['id']}/decision",
                             json={'decision':'PASS','evaluation':mismatch.json()['summary']})
        assert decision.status_code==422
        assert client.get('/api/v1/applications/'+aid).json()['application']['stage']=='ROUND1_SCHEDULED'


def test_transcript_text_limit_is_explicit():
    from app.services.documents import MAX_TRANSCRIPT_CHARS
    try:
        read_transcript(('字'*(MAX_TRANSCRIPT_CHARS+1)).encode(), '.txt')
    except ValueError as exc:
        assert '不得超过' in str(exc)
    else:
        raise AssertionError('Oversized transcripts must be rejected, not truncated')


def test_text_resume_can_be_uploaded_for_assessment_and_rejects_invalid_utf8():
    with TestClient(app) as client:
        text='姓名：林沐安\n岗位：Python后端工程师\n技能：Python、SQL'
        uploaded=client.post('/api/v1/files',files={'file':('resume.txt',text.encode('utf-8'),'text/plain')})
        assert uploaded.status_code==201,uploaded.text
        assessment=client.post('/api/v1/resume-assessments',json={
            'file_id':uploaded.json()['id'],'ai_config':{'mode':'demo'}})
        assert assessment.status_code==201
        assert assessment.json()['evaluation']['candidate_name']=='林沐安'
        invalid=client.post('/api/v1/files',files={'file':('resume.txt',b'\xff\xfe','text/plain')})
        assert invalid.status_code==400


def test_demo_transcript_upload_needs_no_model():
    with TestClient(app) as client:
        created=send(client,'姓名：张晨，岗位：Python后端工程师，这是简历')
        aid=created.json()['application_id']
        send(client,'HR初筛评价：建议面试',aid)
        send(client,'安排一面，2026-11-02 15:00，面试官李工',aid)
        interview=client.get('/api/v1/applications/'+aid).json()['interviews'][0]
        text='姓名：张晨\n岗位：Python后端工程师\n一面\n面试官：介绍项目\n候选人：我做了 Python 接口'
        response=client.post(f"/api/v1/interviews/{interview['id']}/transcripts",
                             data={'ai_mode':'demo'},
                             files={'file':('interview.txt',text.encode(),'text/plain')})
        assert response.status_code==201,response.text
        assert response.json()['model']=='demo'
        assert '非 AI' in response.json()['summary']['summary']
        assert client.get('/api/v1/applications/'+aid).json()['application']['stage']=='ROUND1_SCHEDULED'


def test_custom_model_key_is_not_saved_in_message_or_graph(monkeypatch):
    import json
    from app.agents import proposal as proposal_agent
    from app.db import SessionLocal
    from app.entities import Message

    secret='private-test-key'
    monkeypatch.setattr(proposal_agent,'complete_json',lambda *_args,**_kwargs:
                        '{"intent":"RESUME_RECEIVED","candidate_name":"张晨","position":"Python后端工程师"}')
    with TestClient(app) as client:
        response=client.post('/api/v1/messages',json={
            'external_message_id':'custom-model-message','content':'这是张晨的简历',
            'ai_config':{'mode':'custom','model':'test-model',
                         'base_url':'https://example.com/v1','api_key':secret}})
        assert response.status_code==200,response.text
        message_id=response.json()['message_id']
        state=app.state.graph.get_state({'configurable':{'thread_id':message_id}}).values
        assert secret not in json.dumps(state,ensure_ascii=False)
        with SessionLocal() as db:
            saved=db.get(Message,message_id)
            assert secret not in saved.content and secret not in saved.result


def test_seed_data_and_complete_quick_demo():
    from sqlalchemy import func, select
    from app.db import SessionLocal
    from app.entities import Candidate, Outbox
    from app.services.demo_seed import demo_id, seed_demo_data

    with TestClient(app) as client:
        with SessionLocal() as db:
            candidate=db.get(Candidate,demo_id('candidate:gu'))
            candidate.skills='已编辑技能'
            db.commit()
            count=db.scalar(select(func.count(Candidate.id)))
            jobs=db.scalar(select(func.count(Outbox.id)))
            seed_demo_data(db)
            assert db.scalar(select(func.count(Candidate.id)))==count
            assert db.get(Candidate,candidate.id).skills=='已编辑技能'
            assert db.scalar(select(func.count(Outbox.id)))==jobs
        examples=client.get('/api/v1/demo/examples').json()
        expected=['ROUND1_PENDING','ROUND1_SCHEDULED','ROUND2_PENDING','ROUND2_SCHEDULED',
                  'ROUND3_PENDING','ROUND3_SCHEDULED','OFFER_PENDING']
        for example, stage in zip(examples,expected,strict=True):
            response=send(client,example['text'],example['application_id'])
            assert response.status_code==200,response.text
            assert response.json().get('new_stage')==stage,response.json()


def test_management_tables_create_edit_and_follow_workflow():
    with TestClient(app) as client:
        position=next(p for p in client.get('/api/v1/positions').json() if p['title']=='Python后端工程师')
        response=client.post('/api/v1/management/applications',json={
            'candidate_name':'韩星','position_id':position['id'],'education':'本科',
            'skills':'Python','work_years':3,'owner':'HR一号'})
        assert response.status_code==201,response.text
        application=response.json();aid=application['id']
        edit={'name':'韩星','education':'硕士','skills':'Python、SQL','work_years':4,
              'owner':'HR二号','version':application['version']}
        assert client.put(f'/api/v1/management/applications/{aid}',json=edit).status_code==200
        assert client.put(f'/api/v1/management/applications/{aid}',json=edit).status_code==409
        assert client.get(f'/api/v1/applications/{aid}').json()['candidate']['skills']=='Python、SQL'
        schedule={'action':'schedule','round_no':1,'scheduled_at':'2026-11-02 15:00','interviewer':'李工'}
        assert client.post(f'/api/v1/management/applications/{aid}/interview-actions',json=schedule).status_code==422
        assert send(client,'HR初筛评价：经验匹配，建议面试',aid).json()['new_stage']=='ROUND1_PENDING'
        scheduled=client.post(f'/api/v1/management/applications/{aid}/interview-actions',json=schedule)
        assert scheduled.status_code==200,scheduled.text
        interview=next(i for i in client.get('/api/v1/management/interviews').json() if i['application_id']==aid)
        update={'scheduled_at':'2026-11-03 15:00','interviewer':'王工','evaluation':None}
        assert client.put(f"/api/v1/management/interviews/{interview['id']}",json=update).status_code==200
        feedback={'action':'feedback','round_no':1,'evaluation':'技术扎实，一面通过','decision':'PASS'}
        result=client.post(f'/api/v1/management/applications/{aid}/interview-actions',json=feedback)
        assert result.status_code==200,result.text
        assert result.json()['new_stage']=='ROUND2_PENDING'
        detail=client.get(f'/api/v1/applications/{aid}').json()
        assert detail['interviews'][0]['interviewer']=='王工'
        assert any(e['type']=='MANUAL_INTERVIEW_UPDATED' for e in detail['events'])


def test_manual_candidate_resume_and_interview_workflow():
    with TestClient(app) as client:
        created=client.post('/api/v1/management/applications',json={
            'candidate_name':'沈初','position_title':'数据分析师','phone':'123456789',
            'email':'demo@example.com','skills':'SQL','owner':'HR'})
        assert created.status_code==201,created.text
        aid=created.json()['id']
        manual=client.post(f'/api/v1/management/applications/{aid}/resumes/manual',
                           json={'text':'姓名：沈初\n项目：数据分析'})
        assert manual.status_code==201,manual.text
        document=Document();document.add_paragraph('姓名：沈初');document.add_paragraph('技能：SQL')
        buffer=io.BytesIO();document.save(buffer)
        uploaded=client.post(f'/api/v1/management/applications/{aid}/resumes',
                             files={'file':('resume.docx',buffer.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        assert uploaded.status_code==201,uploaded.text
        detail=client.get(f'/api/v1/applications/{aid}').json()
        assert detail['candidate']['email']=='demo@example.com'
        assert len(detail['resumes'])==2
        assert client.get(f"/api/v1/management/resumes/{manual.json()['id']}/download").content.decode('utf-8').startswith('姓名：沈初')
        screening=client.post(f'/api/v1/management/applications/{aid}/screening',
                              json={'decision':'INTERVIEW','evaluation':'技能匹配，进入一面'})
        assert screening.status_code==200 and screening.json()['stage']=='ROUND1_PENDING'
        assert client.post(f'/api/v1/management/applications/{aid}/screening',
                           json={'decision':'REJECT','evaluation':'重复操作'}).status_code==409
        schedule={'action':'schedule','round_no':1,'scheduled_at':'2026-11-02 15:00','interviewer':'李工'}
        assert client.post(f'/api/v1/management/applications/{aid}/interview-actions',json=schedule).status_code==200
        feedback={'action':'feedback','round_no':1,'evaluation':'项目经历待核实','decision':'PENDING'}
        assert client.post(f'/api/v1/management/applications/{aid}/interview-actions',json=feedback).json()['new_stage']=='ROUND1_FEEDBACK'
        feedback.update(evaluation='核实后表现良好',decision='PASS')
        assert client.post(f'/api/v1/management/applications/{aid}/interview-actions',json=feedback).json()['new_stage']=='ROUND2_PENDING'
        detail=client.get(f'/api/v1/applications/{aid}').json()
        assert detail['interviews'][0]['evaluation']=='核实后表现良好'
        assert any(e['type']=='MANUAL_SCREENING' for e in detail['events'])


def test_candidate_profile_edit_rejects_stale_other_application():
    with TestClient(app) as client:
        first=client.post('/api/v1/management/applications',json={
            'candidate_name':'赵宁','position_title':'前端工程师','owner':'HR'}).json()
        second=client.post('/api/v1/management/applications',json={
            'candidate_id':first['candidate_id'],'position_title':'数据分析师','owner':'HR'}).json()
        edit={'name':'赵宁','education':'本科','skills':'Vue','work_years':2,'owner':'HR','version':first['version']}
        assert client.put(f"/api/v1/management/applications/{first['id']}",json=edit).status_code==200
        edit.update(skills='过时信息',version=second['version'])
        assert client.put(f"/api/v1/management/applications/{second['id']}",json=edit).status_code==409
        assert client.get(f"/api/v1/applications/{second['id']}").json()['candidate']['skills']=='Vue'


def test_resume_assessment_cannot_reassign_file_to_another_candidate():
    with TestClient(app) as client:
        document=Document();document.add_paragraph('姓名：顾清禾');document.add_paragraph('岗位：前端工程师')
        buffer=io.BytesIO();document.save(buffer)
        uploaded=client.post('/api/v1/files',files={'file':('resume.docx',buffer.getvalue(),
                             'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()
        from app.services.demo_seed import demo_id
        # A manually uploaded candidate resume is already owned, even when reused in an assessment.
        from app.db import SessionLocal
        from app.entities import Resume
        with SessionLocal() as db:
            resume=db.get(Resume,uploaded['id']);resume.candidate_id=demo_id('candidate:gu');db.commit()
        assessment=client.post('/api/v1/resume-assessments',json={'file_id':uploaded['id'],'ai_config':{'mode':'demo'}}).json()
        evaluation=assessment['evaluation'];evaluation['candidate_name']='另一位候选人'
        response=client.post(f"/api/v1/resume-assessments/{assessment['id']}/decision",json={'decision':'PENDING','evaluation':evaluation})
        assert response.status_code==422 and '简历已关联其他候选人' in response.json()['detail']
        with SessionLocal() as db:
            assert db.get(Resume,uploaded['id']).candidate_id==demo_id('candidate:gu')


def test_manual_resume_upload_rolls_back_file_and_record_on_failure(monkeypatch):
    from sqlalchemy import func, select
    from app.api import management
    from app.config import UPLOAD_DIR
    from app.db import SessionLocal
    from app.entities import Resume
    from app.services.demo_seed import demo_id

    with TestClient(app, raise_server_exceptions=False) as client:
        with SessionLocal() as db:
            before=db.scalar(select(func.count()).select_from(Resume))
        files_before={path.name for path in UPLOAD_DIR.iterdir()}
        monkeypatch.setattr(management,'record_event',lambda *_args,**_kwargs: (_ for _ in ()).throw(RuntimeError('failed')))
        document=Document();document.add_paragraph('姓名：顾清禾')
        buffer=io.BytesIO();document.save(buffer)
        response=client.post(f"/api/v1/management/applications/{demo_id('application:gu')}/resumes",
                             files={'file':('resume.docx',buffer.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        assert response.status_code==500
        with SessionLocal() as db:
            assert db.scalar(select(func.count()).select_from(Resume))==before
        assert {path.name for path in UPLOAD_DIR.iterdir()}==files_before


def test_concurrent_message_application_is_idempotent():
    from concurrent.futures import ThreadPoolExecutor
    from sqlalchemy import select
    from app.db import SessionLocal
    from app.entities import Event, Message
    from app.schemas import Proposal
    from app.services.demo_seed import demo_id
    from app.services.recruitment import apply_proposal

    with TestClient(app):
        aid=demo_id('application:lin')
        with SessionLocal() as db:
            message=Message(source='web',external_message_id='concurrent-test',content='初筛')
            db.add(message);db.commit();message_id=message.id
        proposal=Proposal(intent='HR_EVALUATION',application_id=aid,evaluation='建议面试')
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda _:apply_proposal(proposal,message_id,'hr-demo'),range(2)))
        assert all(result['status']=='applied' for result in results)
        with SessionLocal() as db:
            assert len(db.scalars(select(Event).where(Event.message_id==message_id)).all())==1


def test_resume_assessment_can_clear_wrong_application():
    with TestClient(app) as client:
        original_id=client.get('/api/v1/demo/examples').json()[0]['application_id']
        before=client.get(f'/api/v1/applications/{original_id}').json()
        document=Document()
        for line in ['姓名：许宁','岗位：Python后端工程师','技能：Python、SQL']:
            document.add_paragraph(line)
        content=io.BytesIO();document.save(content)
        upload=client.post('/api/v1/files',files={
            'file':('resume.docx',content.getvalue(),
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        assert upload.status_code==201,upload.text
        response=client.post('/api/v1/resume-assessments',json={
            'file_id':upload.json()['id'],'application_id':original_id,'ai_config':{'mode':'demo'}})
        assert response.status_code==201,response.text
        path=f"/api/v1/resume-assessments/{response.json()['id']}/decision"
        fields={'decision':'PENDING','evaluation':response.json()['evaluation']}
        unchanged=client.post(path,json={**fields,'application_id':original_id})
        assert unchanged.status_code==422
        assert '简历信息与所选申请不一致' in unchanged.json()['detail']
        cleared=client.post(path,json={**fields,'application_id':None})
        assert cleared.status_code==200,cleared.text
        assert cleared.json()['application_id']!=original_id
        assert client.get(f"/api/v1/applications/{cleared.json()['application_id']}").json()['candidate']['name']=='许宁'
        after=client.get(f'/api/v1/applications/{original_id}').json()
        assert after['candidate']['name']==before['candidate']['name']
        assert after['application']['version']==before['application']['version']


def test_existing_resume_review_can_still_clear_wrong_application():
    import json
    from app.db import SessionLocal
    from app.entities import Message, Review

    with TestClient(app) as client:
        original_id=client.get('/api/v1/demo/examples').json()[0]['application_id']
        with SessionLocal() as db:
            msg=Message(source='web',external_message_id='old-resume-review',content='这是许宁的简历')
            db.add(msg);db.commit();message_id=msg.id
        thread={'configurable':{'thread_id':message_id}}
        state={'content':'这是许宁的简历','resume_text':'姓名：许宁\n岗位：Python后端工程师','application_id':original_id}
        interrupted=app.state.graph.invoke(state,config=thread)
        assert '__interrupt__' in interrupted
        values=app.state.graph.get_state(thread).values
        with SessionLocal() as db:
            review=Review(message_id=message_id,application_id=original_id,
                          proposal=json.dumps(values['proposal'],ensure_ascii=False),
                          reason=values['reason'],thread_id=message_id)
            db.add(review);db.commit();review_id=review.id
        resolved=client.post(f'/api/v1/reviews/{review_id}/resolve',json={
            'action':'approve','application_id':None,'candidate_name':'许宁','position':'Python后端工程师'})
        assert resolved.status_code==200,resolved.text
        assert resolved.json()['application_id']!=original_id


def test_model_null_decision_does_not_auto_pass_interview(monkeypatch):
    from app.agents import proposal as proposal_agent
    from app.services.demo_seed import demo_id

    monkeypatch.setattr(proposal_agent,'complete_json',lambda *_args,**_kwargs:
                        '{"intent":"INTERVIEW_FEEDBACK","round_no":1,'
                        '"evaluation":"讨论了前端项目经验","decision":null}')
    with TestClient(app) as client:
        app_id=demo_id('application:gu')
        response=client.post('/api/v1/messages',json={
            'external_message_id':'null-decision-test','application_id':app_id,
            'content':'一面结束，讨论了前端项目经验',
            'ai_config':{'mode':'custom','model':'test-model',
                         'base_url':'https://example.com/v1','api_key':'test-key'}})
        assert response.status_code==200,response.text
        assert response.json()['new_stage']=='ROUND1_FEEDBACK'
        detail=client.get(f'/api/v1/applications/{app_id}').json()
        assert detail['interviews'][0]['result']=='PENDING'


def test_resume_assessment_waits_for_editable_human_decision():
    with TestClient(app) as client:
        doc=Document()
        for line in ('姓名：林初','岗位：Python后端工程师','技能：Python'):
            doc.add_paragraph(line)
        content=io.BytesIO();doc.save(content)
        upload=client.post('/api/v1/files',files={'file':('resume.docx',content.getvalue(),
                           'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        file_id=upload.json()['id']
        assert send(client,'简历',file_id=file_id).status_code==422
        generated=client.post('/api/v1/resume-assessments',json={'file_id':file_id,'ai_config':{'mode':'demo'}})
        assert generated.status_code==201,generated.text
        assessment=generated.json()
        assert assessment['status']=='PENDING' and assessment['model']=='demo'
        assert not any(a['candidate_name']=='林初' for a in client.get('/api/v1/applications').json())
        evaluation=assessment['evaluation'];evaluation['summary']='面试官修改后的评价'
        evaluation['skills']='Python、SQL'
        path=f"/api/v1/resume-assessments/{assessment['id']}/decision"
        pending=client.post(path,json={'decision':'PENDING','evaluation':evaluation})
        assert pending.status_code==200,pending.text
        aid=pending.json()['application_id']
        assert pending.json()['new_stage']=='SCREENING'
        saved=next(x for x in client.get('/api/v1/resume-assessments').json() if x['id']==assessment['id'])
        assert saved['status']=='DEFERRED' and saved['reviewed']['summary']=='面试官修改后的评价'
        final=client.post(path,json={'decision':'INTERVIEW','application_id':aid,'evaluation':evaluation})
        assert final.status_code==200,final.text
        assert final.json()['new_stage']=='ROUND1_PENDING'
        assert client.post(path,json={'decision':'INTERVIEW','application_id':aid,'evaluation':evaluation}).status_code==422
        detail=client.get(f'/api/v1/applications/{aid}').json()
        assert detail['candidate']['skills']=='Python、SQL'
        assert len([e for e in detail['events'] if e['type']=='RESUME_ASSESSMENT_CONFIRMED'])==2


def test_interview_assessment_waits_for_human_and_allows_correction():
    from app.services.demo_seed import demo_id
    with TestClient(app) as client:
        aid=demo_id('application:gu')
        interview=client.get(f'/api/v1/applications/{aid}').json()['interviews'][0]
        text='姓名：其他人\n一面\n面试官：介绍项目\n候选人：完成 Vue 项目'
        uploaded=client.post(f"/api/v1/interviews/{interview['id']}/transcripts",data={'ai_mode':'demo'},
                             files={'file':('interview.txt',text.encode(),'text/plain')})
        assert uploaded.status_code==201,uploaded.text
        transcript=uploaded.json();path=f"/api/v1/interviews/{interview['id']}/transcripts/{transcript['id']}/decision"
        evaluation=transcript['summary']
        assert client.get(f'/api/v1/applications/{aid}').json()['application']['stage']=='ROUND1_SCHEDULED'
        assert client.post(path,json={'decision':'PASS','evaluation':evaluation}).status_code==422
        evaluation['candidate_name']='顾清禾';evaluation['summary']='面试官修正后的面试总结';evaluation['interviewer']='王工'
        evaluation['performance']=['描述了 Vue 项目']
        evaluation['scores']=[{'dimension':'项目实践','score':4,'evidence':'完成 Vue 项目','reasoning':'有相关经验'}]
        pending=client.post(path,json={'decision':'PENDING','evaluation':evaluation})
        assert pending.status_code==200,pending.text
        assert pending.json()['new_stage']=='ROUND1_FEEDBACK'
        detail=client.get(f'/api/v1/applications/{aid}').json()
        assert detail['interviews'][0]['transcript']['review']['evaluation']['summary']=='面试官修正后的面试总结'
        assert detail['interviews'][0]['transcript']['review']['evaluation']['scores'][0]['score']==4
        assert detail['interviews'][0]['transcript']['review']['evaluation']['performance']==['描述了 Vue 项目']
        assert detail['interviews'][0]['interviewer']=='王工'
        after_completion=client.post(f"/api/v1/interviews/{interview['id']}/transcripts",data={'ai_mode':'demo'},
                                     files={'file':('late.txt',b'late text','text/plain')})
        assert after_completion.status_code==409
        passed=client.post(path,json={'decision':'PASS','evaluation':evaluation})
        assert passed.status_code==200,passed.text
        assert passed.json()['new_stage']=='ROUND2_PENDING'
        assert client.post(path,json={'decision':'PASS','evaluation':evaluation}).status_code==422


def test_manual_interview_feedback_finalizes_deferred_assessment():
    from app.services.demo_seed import demo_id
    with TestClient(app) as client:
        aid=demo_id('application:gu')
        interview=client.get(f'/api/v1/applications/{aid}').json()['interviews'][0]
        text='候选人姓名：顾清禾\n应聘岗位：前端工程师\n第一轮\n候选人：完成 Vue 项目'
        uploaded=client.post(f"/api/v1/interviews/{interview['id']}/transcripts",data={'ai_mode':'demo'},
                             files={'file':('interview.txt',text.encode(),'text/plain')}).json()
        path=f"/api/v1/interviews/{interview['id']}/transcripts/{uploaded['id']}/decision"
        evaluation=uploaded['summary'];evaluation['summary']='首次面试总结'
        assert client.post(path,json={'decision':'PENDING','evaluation':evaluation}).status_code==200
        feedback=client.post(f'/api/v1/management/applications/{aid}/interview-actions',json={
            'action':'feedback','round_no':1,'decision':'PASS','evaluation':'补充核实后通过'})
        assert feedback.status_code==200 and feedback.json()['new_stage']=='ROUND2_PENDING'
        iv=client.get(f'/api/v1/applications/{aid}').json()['interviews'][0]
        assert iv['result']=='PASS' and iv['evaluation']=='补充核实后通过'
        assert iv['transcript']['review']['decision']=='PASS'
        assert iv['transcript']['review']['status']=='FINAL'
        assert iv['transcript']['review']['evaluation']['summary']=='首次面试总结'
        assert client.post(path,json={'decision':'PASS','evaluation':evaluation}).status_code==422


def test_manual_offer_and_hire_actions_validate_stage_and_record_events():
    from app.db import SessionLocal
    from app.entities import Application
    from app.services.demo_seed import demo_id
    with TestClient(app) as client:
        aid=demo_id('application:gu')
        path=f'/api/v1/management/applications/{aid}/employment-action'
        assert client.post(path,json={'target':'HIRED'}).status_code==422
        with SessionLocal() as db:
            candidate_app=db.get(Application,aid);candidate_app.stage='OFFER_PENDING';db.commit()
        accepted=client.post(path,json={'target':'OFFER_ACCEPTED'})
        assert accepted.status_code==200 and accepted.json()['new_stage']=='OFFER_ACCEPTED'
        assert client.post(path,json={'target':'OFFER_ACCEPTED'}).status_code==422
        hired=client.post(path,json={'target':'HIRED'})
        assert hired.status_code==200 and hired.json()['new_stage']=='HIRED'
        detail=client.get(f'/api/v1/applications/{aid}').json()
        assert [e['type'] for e in detail['events']].count('STATUS_UPDATE')==2


def test_complete_recruitment_with_manual_followup_and_employment():
    with TestClient(app) as client:
        document=Document()
        for line in ('姓名：苏知行','岗位：前端工程师','技能：Vue、TypeScript'):
            document.add_paragraph(line)
        buffer=io.BytesIO();document.save(buffer)
        uploaded=client.post('/api/v1/files',files={'file':('resume.docx',buffer.getvalue(),
                             'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()
        assessment=client.post('/api/v1/resume-assessments',json={
            'file_id':uploaded['id'],'ai_config':{'mode':'demo'}}).json()
        created=client.post(f"/api/v1/resume-assessments/{assessment['id']}/decision",json={
            'decision':'INTERVIEW','evaluation':assessment['evaluation']})
        assert created.status_code==200,created.text
        aid=created.json()['application_id']
        for round_no in (1,2,3):
            scheduled=client.post(f'/api/v1/management/applications/{aid}/interview-actions',json={
                'action':'schedule','round_no':round_no,'scheduled_at':f'2026-11-0{round_no+1} 15:00','interviewer':'李工'})
            assert scheduled.status_code==200,scheduled.text
            iv=next(i for i in client.get(f'/api/v1/applications/{aid}').json()['interviews'] if i['round_no']==round_no)
            text=(f'候选人姓名：苏知行\n应聘岗位：前端工程师\n面试官姓名：李工\n第{round_no}轮\n'
                  '李工：请介绍项目\n苏知行：使用 Vue 开发页面')
            transcript=client.post(f"/api/v1/interviews/{iv['id']}/transcripts",data={'ai_mode':'demo'},
                                   files={'file':('interview.txt',text.encode(),'text/plain')})
            assert transcript.status_code==201,transcript.text
            summary=transcript.json()['summary']
            assert summary['performance']==['候选人原话：使用 Vue 开发页面']
            path=f"/api/v1/interviews/{iv['id']}/transcripts/{transcript.json()['id']}/decision"
            if round_no==1:
                assert client.post(path,json={'decision':'PENDING','evaluation':summary}).status_code==200
                result=client.post(f'/api/v1/management/applications/{aid}/interview-actions',json={
                    'action':'feedback','round_no':1,'decision':'PASS','evaluation':'人工确认推进'})
            else:
                result=client.post(path,json={'decision':'PASS','evaluation':summary})
            expected=f'ROUND{round_no+1}_PENDING' if round_no<3 else 'OFFER_PENDING'
            assert result.status_code==200 and result.json()['new_stage']==expected
        employment=f'/api/v1/management/applications/{aid}/employment-action'
        assert client.post(employment,json={'target':'OFFER_ACCEPTED'}).json()['new_stage']=='OFFER_ACCEPTED'
        assert client.post(employment,json={'target':'HIRED'}).json()['new_stage']=='HIRED'
        detail=client.get(f'/api/v1/applications/{aid}').json()
        assert len(detail['interviews'])==3
        assert all(i['transcript']['review']['decision']=='PASS' for i in detail['interviews'])


def test_resume_agent_uses_config_and_retains_human_rejection(monkeypatch):
    import json
    from app.agents import resume_assessment as agent
    captured={}

    def fake_model(system, prompt, timeout, config):
        captured.update(system=system, prompt=prompt, timeout=timeout, mode=config.mode)
        return json.dumps({'candidate_name':'周夏','position':'Python后端工程师',
                           'summary':'候选人描述了接口开发经验',
                           'strengths':['有相关项目经历'],'suggestions':['核实接口设计经验']},ensure_ascii=False)

    monkeypatch.setattr(agent,'complete_json',fake_model)
    with TestClient(app) as client:
        document=Document();document.add_paragraph('姓名：周夏');document.add_paragraph('岗位：Python后端工程师')
        content=io.BytesIO();document.save(content)
        upload=client.post('/api/v1/files',files={'file':('resume.docx',content.getvalue(),
                           'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        created=client.post('/api/v1/resume-assessments',json={
            'file_id':upload.json()['id'],'context':'请作为面试参考',
            'ai_config':{'mode':'custom','model':'test-model','base_url':'https://example.com/v1','api_key':'private-test-key'}})
        assert created.status_code==201,created.text
        assert captured['mode']=='custom' and captured['timeout']==90
        assert '姓名：周夏' in captured['prompt'] and '请作为面试参考' in captured['prompt']
        assert 'private-test-key' not in created.text
        evaluation=created.json()['evaluation'];evaluation['summary']='面试官确认的评价'
        result=client.post(f"/api/v1/resume-assessments/{created.json()['id']}/decision",
                           json={'decision':'REJECT','evaluation':evaluation})
        assert result.status_code==200,result.text
        assert result.json()['new_stage']=='REJECTED'
        record=next(r for r in client.get('/api/v1/resume-assessments').json() if r['id']==created.json()['id'])
        assert record['status']=='FINAL' and record['reviewed']['summary']=='面试官确认的评价'
        assert 'private-test-key' not in json.dumps(record,ensure_ascii=False)


def test_only_latest_interview_transcript_can_be_confirmed():
    from app.services.demo_seed import demo_id
    with TestClient(app) as client:
        aid=demo_id('application:gu')
        interview=client.get(f'/api/v1/applications/{aid}').json()['interviews'][0]
        base=f"/api/v1/interviews/{interview['id']}/transcripts"
        old=client.post(base,data={'ai_mode':'demo'},files={'file':('old.txt',b'old text','text/plain')}).json()
        new=client.post(base,data={'ai_mode':'demo'},files={'file':('new.txt',b'new text','text/plain')}).json()
        stale=client.post(f"{base}/{old['id']}/decision",json={'decision':'REJECT','evaluation':old['summary']})
        assert stale.status_code==422 and '更新的' in stale.json()['detail']
        assert client.get(f'/api/v1/applications/{aid}').json()['application']['stage']=='ROUND1_SCHEDULED'
        rejected=client.post(f"{base}/{new['id']}/decision",json={'decision':'REJECT','evaluation':new['summary']})
        assert rejected.status_code==200,rejected.text
        assert rejected.json()['new_stage']=='REJECTED'


def test_transcript_upload_rechecks_stage_after_model_processing(monkeypatch):
    from app.api import interviews as api
    from app.db import SessionLocal
    from app.entities import Application, Interview
    from app.services.demo_seed import demo_id

    aid=demo_id('application:gu')
    with TestClient(app) as client:
        interview=client.get(f'/api/v1/applications/{aid}').json()['interviews'][0]
        def finish_interview(*_args):
            with SessionLocal() as db:
                db.get(Application,aid).stage='ROUND1_FEEDBACK'
                db.get(Interview,interview['id']).status='COMPLETED'
                db.commit()
            return InterviewSummary(summary='测试总结')
        monkeypatch.setattr(api,'summarize_interview',finish_interview)
        response=client.post(f"/api/v1/interviews/{interview['id']}/transcripts",
                             data={'ai_mode':'demo'},files={'file':('test.txt',b'test','text/plain')})
        assert response.status_code==409,response.text
        assert client.get(f'/api/v1/applications/{aid}').json()['interviews'][0]['transcript'] is None


def test_local_resume_model_common_json_variants(monkeypatch):
    from app.agents import resume_assessment as agent
    raw='''<think>internal text</think>
```json
{"candidate_name":"陈安", "position":"前端工程师", "skills":["Vue","TypeScript"],
 "work_years":"3年", "summary":["负责前端项目", "使用 Vue"],
 "strengths":null, "concerns":[], "suggestions":"核实项目职责"}
```'''
    monkeypatch.setattr(agent,'complete_json',lambda *_args,**_kwargs:raw)
    with TestClient(app) as client:
        document=Document();document.add_paragraph('姓名：陈安')
        content=io.BytesIO();document.save(content)
        upload=client.post('/api/v1/files',files={'file':('resume.docx',content.getvalue(),
                           'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        response=client.post('/api/v1/resume-assessments',json={'file_id':upload.json()['id'],
            'ai_config':{'mode':'custom','model':'local','base_url':'http://localhost:11434/v1'}})
        assert response.status_code==201,response.text
        evaluation=response.json()['evaluation']
        assert evaluation['skills']=='Vue、TypeScript' and evaluation['work_years']==3
        assert evaluation['summary']=='负责前端项目\n使用 Vue'
        assert evaluation['strengths']==[] and evaluation['suggestions']==['核实项目职责']


def test_resume_model_errors_are_useful_and_do_not_expose_content(monkeypatch, caplog):
    from app.agents import resume_assessment as agent
    import httpx
    from openai import APIConnectionError
    with TestClient(app) as client:
        document=Document();document.add_paragraph('姓名：测试候选人')
        content=io.BytesIO();document.save(content)
        upload=client.post('/api/v1/files',files={'file':('resume.docx',content.getvalue(),
                           'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        payload={'file_id':upload.json()['id'],'ai_config':{'mode':'custom','model':'local',
                 'base_url':'http://localhost:11434/v1','api_key':'private-test-key'}}
        monkeypatch.setattr(agent,'complete_json',lambda *_args,**_kwargs:
                            '{"summary":{"sensitive":"private-test-key"}}')
        invalid=client.post('/api/v1/resume-assessments',json=payload)
        assert invalid.status_code==502 and 'summary' in invalid.json()['detail']
        assert 'private-test-key' not in invalid.text and 'private-test-key' not in caplog.text
        def fail(*_args,**_kwargs):
            raise APIConnectionError(request=httpx.Request('POST','http://localhost:11434/v1'))
        monkeypatch.setattr(agent,'complete_json',fail)
        disconnected=client.post('/api/v1/resume-assessments',json=payload)
        assert disconnected.status_code==502 and '无法连接模型服务' in disconnected.json()['detail']
